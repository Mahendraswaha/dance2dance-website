import openpyxl
import json

file_path = r'C:\Renas\Antigravity\Dance2Dance_Traducoes_Revisao.xlsx'

# Flatten JSON helper
def flatten_json(y):
    out = {}
    def flatten(x, name=''):
        if type(x) is dict:
            for a in x:
                flatten(x[a], name + a + '.')
        elif type(x) is list:
            for i, a in enumerate(x):
                flatten(a, name + str(i) + '.')
        else:
            out[name[:-1]] = x
    flatten(y)
    return out

pt = flatten_json(json.load(open('src/i18n/locales/pt.json', encoding='utf-8')))
en = flatten_json(json.load(open('src/i18n/locales/en.json', encoding='utf-8')))
no = flatten_json(json.load(open('src/i18n/locales/no.json', encoding='utf-8')))

# Load workbook safely to preserve all formatting
wb = openpyxl.load_workbook(file_path)
ws = wb.active

# Iterate through rows and update values only
for row in range(2, ws.max_row + 1):
    chave = ws.cell(row=row, column=3).value # Column C is 'Chave Interna'
    
    if chave:
        if chave in pt:
            ws.cell(row=row, column=4).value = pt[chave] # Column D is PT
        if chave in en:
            ws.cell(row=row, column=5).value = en[chave] # Column E is EN
        if chave in no:
            ws.cell(row=row, column=6).value = no[chave] # Column F is NO

wb.save(file_path)
print('Excel synced successfully using openpyxl (formatting preserved).')
