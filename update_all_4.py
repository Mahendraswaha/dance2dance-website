import json
import openpyxl

# Update pt.json
pt_path = "C:/Renas/Antigravity/Website-builder/src/i18n/locales/pt.json"
with open(pt_path, 'r', encoding='utf-8') as f:
    pt_data = json.load(f)

pt_quote = "\"Quero que <i>Good Morning Dance</i> seja um presente que você dá a si mesmo antes do dia começar. Apenas um pouco de música, alguns movimentos e uma oportunidade de se sentir bem no seu próprio corpo.\""
pt_data["good_morning_dance"]["different_way"]["closingQuote"] = pt_quote

with open(pt_path, 'w', encoding='utf-8') as f:
    json.dump(pt_data, f, ensure_ascii=False, indent=2)

# Update en.json
en_path = "C:/Renas/Antigravity/Website-builder/src/i18n/locales/en.json"
with open(en_path, 'r', encoding='utf-8') as f:
    en_data = json.load(f)

en_quote = "\"I want <i>Good Morning Dance</i> to be a little gift that you give yourself before the day begins. Just some music, some movement, and an opportunity to feel good in your own body.\""
en_data["good_morning_dance"]["different_way"]["closingQuote"] = en_quote

with open(en_path, 'w', encoding='utf-8') as f:
    json.dump(en_data, f, ensure_ascii=False, indent=2)

# Update no.json
no_path = "C:/Renas/Antigravity/Website-builder/src/i18n/locales/no.json"
with open(no_path, 'r', encoding='utf-8') as f:
    no_data = json.load(f)

no_quote = "\"Jeg vil at <i>Good Morning Dance</i> skal være en liten gave som du gir deg selv før dagen begynner. Bare litt musikk, litt bevegelse og en mulighet til å føle deg vel i din egen kropp.\""
no_data["good_morning_dance"]["different_way"]["closingQuote"] = no_quote

with open(no_path, 'w', encoding='utf-8') as f:
    json.dump(no_data, f, ensure_ascii=False, indent=2)

# Update Excel
wb = openpyxl.load_workbook(r'C:\Renas\Antigravity\Dance2Dance_Traducoes_Revisao.xlsx')
ws = wb.active

for row in range(2, ws.max_row + 1):
    key = ws.cell(row=row, column=3).value
    if key == "good_morning_dance.different_way.closingQuote":
        ws.cell(row=row, column=4).value = pt_quote
        ws.cell(row=row, column=5).value = en_quote
        ws.cell(row=row, column=6).value = no_quote

wb.save(r'C:\Renas\Antigravity\Dance2Dance_Traducoes_Revisao.xlsx')
print("JSONs and Excel updated successfully.")
