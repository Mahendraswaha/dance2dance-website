import json
import openpyxl

# Update pt.json
pt_path = "C:/Renas/Antigravity/Website-builder/src/i18n/locales/pt.json"
with open(pt_path, 'r', encoding='utf-8') as f:
    pt_data = json.load(f)

pt_text = "Biostretch não é pensado como uma aula isolada, mas como um processo contínuo de investigação."
pt_data["regular_classes"]["block2"]["p1"] = pt_text

with open(pt_path, 'w', encoding='utf-8') as f:
    json.dump(pt_data, f, ensure_ascii=False, indent=2)

# Update en.json
en_path = "C:/Renas/Antigravity/Website-builder/src/i18n/locales/en.json"
with open(en_path, 'r', encoding='utf-8') as f:
    en_data = json.load(f)

en_text = "Biostretch is not conceived as an isolated class, but as an ongoing process of investigation."
en_data["regular_classes"]["block2"]["p1"] = en_text

with open(en_path, 'w', encoding='utf-8') as f:
    json.dump(en_data, f, ensure_ascii=False, indent=2)

# Update no.json
no_path = "C:/Renas/Antigravity/Website-builder/src/i18n/locales/no.json"
with open(no_path, 'r', encoding='utf-8') as f:
    no_data = json.load(f)

no_text = "Biostretch er ikke tenkt som en enkeltstående time, men som en kontinuerlig utforskningsprosess."
no_data["regular_classes"]["block2"]["p1"] = no_text

with open(no_path, 'w', encoding='utf-8') as f:
    json.dump(no_data, f, ensure_ascii=False, indent=2)

# Update Excel
wb = openpyxl.load_workbook(r'C:\Renas\Antigravity\Dance2Dance_Traducoes_Revisao.xlsx')
ws = wb.active

for row in range(2, ws.max_row + 1):
    key = ws.cell(row=row, column=3).value
    if key == "regular_classes.block2.p1":
        ws.cell(row=row, column=4).value = pt_text
        ws.cell(row=row, column=5).value = en_text
        ws.cell(row=row, column=6).value = no_text

wb.save(r'C:\Renas\Antigravity\Dance2Dance_Traducoes_Revisao.xlsx')
print("JSONs and Excel updated successfully.")
