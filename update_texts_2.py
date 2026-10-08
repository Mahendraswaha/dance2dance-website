import json
import openpyxl

# Update JSON
filepath = "C:/Renas/Antigravity/Website-builder/src/i18n/locales/pt.json"
with open(filepath, 'r', encoding='utf-8') as f:
    data = json.load(f)

data["good_morning_dance"]["subtitle"] = "Acorde. Mova-se. Divirta-se."
data["good_morning_dance"]["different_way"]["closingQuote"] = "Quero que <i>Good Morning Dance</i> seja um presente que você dá a si mesmo antes do dia começar. Não mais uma coisa que você precisa realizar. Apenas um pouco de música, alguns movimentos e uma oportunidade de se sentir bem no seu próprio corpo."

with open(filepath, 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=2)

# Update Excel
wb = openpyxl.load_workbook(r'C:\Renas\Antigravity\Dance2Dance_Traducoes_Revisao.xlsx')
ws = wb.active

for row in range(2, ws.max_row + 1):
    if ws.cell(row=row, column=3).value == "good_morning_dance.subtitle":
        ws.cell(row=row, column=4).value = "Acorde. Mova-se. Divirta-se."
    if ws.cell(row=row, column=3).value == "good_morning_dance.different_way.closingQuote":
        ws.cell(row=row, column=4).value = "Quero que <i>Good Morning Dance</i> seja um presente que você dá a si mesmo antes do dia começar. Não mais uma coisa que você precisa realizar. Apenas um pouco de música, alguns movimentos e uma oportunidade de se sentir bem no seu próprio corpo."

wb.save(r'C:\Renas\Antigravity\Dance2Dance_Traducoes_Revisao.xlsx')
print("Text updated in pt.json and Excel.")
